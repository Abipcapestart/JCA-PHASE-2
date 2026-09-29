"""Reconstructs a Phase1Output-shaped dict from a jca_phase1 UI/API `.xlsx`
run export (filename pattern `JCA_Phase1_RUN-<timestamp>-<hash>.xlsx`), so it
can be fed into `adapter.adapt_orchestrator_result()` like any other jca_phase1
run - for the (rare, recovery-only) case where the original JSON was never
saved and the API server that produced it is no longer running (its
`RunStore` is in-process/non-persistent - see jca_phase1/CLAUDE.md).

This export format is a REPORTING view, not the full typed output, and is
missing two things Phase 2 needs that this module can never recover on its
own - both must be supplied by the caller, never guessed:

  1. The intervention/product name - no sheet in this export carries it.
  2. `comparator_scenario` per comparator - the "Comparators" sheet has no
     such column (unlike the underlying `ConsolidatedComparator` dataclass).
     Without it, `adapter._adapt_comparator()` correctly skips every
     per-Member-State row (it never guesses a scenario) - so a Phase 2 run
     built from this file alone will have comparators with NO per-state
     requirements, and therefore no real PICO sets. That is a genuine gap in
     THIS xlsx, not a bug in the adapter - fix it upstream (add the column to
     whatever exports this file) for anything beyond a plumbing check.

If you have the original `Phase1Output.to_dict()` JSON, use that with
`adapter.adapt_orchestrator_result()` directly instead - it has no such gaps.
"""

from __future__ import annotations

import hashlib
from collections import defaultdict
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from p2_config import EU_27_MEMBER_STATES


class XlsxRunAdapterError(Exception):
    """Raised for structural problems with the workbook itself (missing
    sheet, run not complete) - never for a single row, which callers should
    expect to see reflected as sparse/empty fields instead."""


def _source_id(url: str) -> str:
    return "src-" + hashlib.sha1((url or "").encode("utf-8")).hexdigest()[:12]


def _sheet_rows(ws) -> List[Dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]


def load_phase1_xlsx(
    path: str, product_name: str, population_label: str = "licensed",
    indication_disease: str = "",
) -> Dict[str, Any]:
    """`product_name` is mandatory and never inferred - see module docstring.
    `indication_disease` defaults to the most common comparator-level
    `indication` text actually present in the "Comparators" sheet (real
    extracted text, not invented) when not supplied."""
    if not product_name:
        raise XlsxRunAdapterError(
            "product_name is required - this xlsx export has no intervention/product name "
            "in any sheet, so it cannot be inferred.")

    wb = load_workbook(path, data_only=True, read_only=True)
    required_sheets = {"Run Summary", "Comparators", "Per-State Verdicts", "Outcomes"}
    missing = required_sheets - set(wb.sheetnames)
    if missing:
        raise XlsxRunAdapterError(
            f"{path!r} is missing expected sheet(s) {sorted(missing)} - not a jca_phase1 run export "
            f"(found: {wb.sheetnames}).")

    summary = {r.get("field"): r.get("value") for r in _sheet_rows(wb["Run Summary"])}
    status = summary.get("status")
    if status and status != "complete":
        raise XlsxRunAdapterError(f"Run status is {status!r}, not 'complete' - nothing usable to adapt.")

    comparator_rows = _sheet_rows(wb["Comparators"])
    verdict_rows = _sheet_rows(wb["Per-State Verdicts"])
    outcome_rows = _sheet_rows(wb["Outcomes"])

    verdicts_by_comparator: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for v in verdict_rows:
        verdicts_by_comparator[v.get("generic_name", "")].append(v)

    comparators = []
    identified_states = set()
    for row in comparator_rows:
        generic_name = row.get("generic_name", "")
        my_verdicts = verdicts_by_comparator.get(generic_name, [])

        sources_by_id: Dict[str, Dict[str, Any]] = {}
        evidence: List[Dict[str, Any]] = []
        per_member_state = []
        for v in my_verdicts:
            url = v.get("source_url") or ""
            sid = _source_id(url) if url else ""
            per_member_state.append({
                "member_state": v.get("member_state", ""),
                "verdict": v.get("verdict", ""),
                "tier": v.get("tier"),
                "source_id": sid,
                "source_url": url,
                "evidence_quote": v.get("evidence_quote", ""),
                "reason": v.get("reason", ""),
            })
            if v.get("verdict") == "standard_of_care":
                identified_states.add(v.get("member_state", ""))
            if sid and sid not in sources_by_id:
                sources_by_id[sid] = {"source_id": sid, "url": url}
            if sid and v.get("evidence_quote"):
                evidence.append({"source_id": sid, "quote": v.get("evidence_quote", "")})

        member_states_field = row.get("member_states") or ""
        tiers = sorted({v.get("tier") for v in my_verdicts
                        if v.get("verdict") == "standard_of_care" and v.get("tier") is not None})

        comparators.append({
            "generic_name": generic_name,
            "brand_names": [],
            "member_states": [s.strip() for s in member_states_field.split(";") if s.strip()],
            "per_member_state": per_member_state,
            "class_or_mechanism": row.get("class_or_mechanism") or "",
            "line_of_therapy": row.get("line_of_therapy") or "",
            "indication": row.get("indication") or "",
            "tiers": tiers,
            "rationale": row.get("rationale") or "",
            # NOT present in this export - see module docstring. Left blank
            # on purpose so the adapter's existing "skip with a warning"
            # path handles it, rather than silently defaulting to a value
            # that would drive the 4-case methodology on a guess.
            "comparator_scenario": "",
            "retain_all_status": row.get("retain_all_status") or "unconfirmed",
            "origin": row.get("origin") or "agent_found",
            "sources": list(sources_by_id.values()),
            "evidence": evidence,
        })

    outcomes: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in outcome_rows:
        bucket = row.get("bucket", "")
        outcomes[bucket].append({
            "concept": row.get("concept", ""),
            "tiers": [],
            "rationale": row.get("rationale") or "",
            "sources": [],
            "evidence": [],
            "listed": bool(row.get("listed", True)),
            "origin": "agent_found",
        })

    if not indication_disease:
        texts = [c["indication"] for c in comparators if c["indication"]]
        indication_disease = max(set(texts), key=texts.count) if texts else ""

    not_identified_states = sorted(set(EU_27_MEMBER_STATES) - identified_states)

    return {
        "input_validation": {"overall_status": "PASS", "fix_items": [], "check_items": []},
        "intervention": {
            "fields": {"product_name_inn": {"value": product_name, "provenance": "confirmed"}},
            "custom_fields": [],
        },
        "populations": [{
            "population_id": population_label,
            "fields": {"indication_disease": {"value": indication_disease, "provenance": "confirmed"}},
            "custom_fields": [],
        }],
        "comparators": comparators,
        "not_identified_states": not_identified_states,
        "member_state_summary": {
            "identified_count": len(identified_states),
            "not_identified_count": 27 - len(identified_states),
            "total": 27,
        },
        "outcomes": {
            "clinical_effectiveness": outcomes.get("clinical_effectiveness", []),
            "safety": outcomes.get("safety", []),
            "quality_of_life": outcomes.get("quality_of_life", []),
            "clinician_patient_reported": outcomes.get("clinician_patient_reported", []),
        },
    }
