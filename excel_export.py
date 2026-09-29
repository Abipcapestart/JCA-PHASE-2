"""Multi-sheet Excel export (task Section 25/26).

Every SME-required output field, for every agent, across separate logical
sheets, with full evidence/provenance traceability - not just the final
PICO sets. Uses openpyxl (already a Phase 1 dependency - build_data_team_
workbook.py / build_validation_workbook.py already use it; no new library
introduced).
"""

from __future__ import annotations

from typing import Any, List

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from schemas import Phase2Input, Phase2RunResult

_HEADER_FONT = Font(bold=True)


def _write_sheet(ws: Worksheet, headers: List[str], rows: List[List[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append([_stringify(v) for v in row])
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, header in enumerate(headers, start=1):
        max_len = max([len(header)] + [len(_stringify(r[i - 1])) for r in rows]) if rows else len(header)
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 60)


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "; ".join(_stringify(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}={_stringify(v)}" for k, v in value.items())
    return str(value)


def build_workbook(result: Phase2RunResult, phase2_input: Phase2Input) -> Workbook:
    wb = Workbook()

    # Sheet 1 - Run Summary
    ws = wb.active
    ws.title = "1. Run Summary"
    total_pico_sets = sum(len(p.assembly.pico_sets) for p in result.populations)
    total_ms = len({ms for p in result.populations for s in p.assembly.pico_sets for ms in s.member_states})
    total_errors_warnings = len(result.logs)
    _write_sheet(ws, ["Field", "Value"], [
        ["run_id", result.run_id],
        ["model", result.model],
        ["overall_status", result.overall_status],
        ["prompt_versions", result.prompt_versions],
        ["population_count", len(result.populations)],
        ["pico_set_count", total_pico_sets],
        ["distinct_member_states_covered", total_ms],
        ["errors_warnings_count", total_errors_warnings],
        ["source_phase1_run_ref", phase2_input.source_phase1_run_ref],
    ])

    # Sheet 2 - Agent 1 Methodology
    ws = wb.create_sheet("2. Agent1 Methodology")
    _write_sheet(ws, ["Field", "Value"], [
        ["methodology_loaded", result.context_locking.methodology_loaded],
        ["guidance_version", result.context_locking.guidance_version],
        ["sections_covered", result.context_locking.sections_covered],
        ["prompt_version", result.prompt_versions.get("context_locking")],
    ])

    # Sheet 3 - Agent 2 Consolidation
    ws = wb.create_sheet("3. Agent2 Consolidation")
    rows = []
    for pop in result.populations:
        for sel in pop.consolidation.selected_comparators:
            tie = sel.tie_break
            rows.append([
                pop.population, sel.generic_name, sel.combination_type, sel.selection_basis,
                sel.member_states, sel.tiers, sel.droppable_note or "",
                sel.source_comparator_refs,
                tie.tie_existed if tie else False,
                tie.reason if tie else "",
            ])
    _write_sheet(ws, ["Population", "Comparator", "Combination Type", "Selection Basis",
                       "Member States", "Tiers", "Droppable Note", "Source Comparator Refs",
                       "Tie-Break Occurred", "Tie-Break Reason"], rows)

    # Sheet 4 - Agent 2 Member State Mapping
    ws = wb.create_sheet("4. Agent2 MS Mapping")
    rows = []
    for pop in result.populations:
        for sel in pop.consolidation.selected_comparators:
            for ms in sel.member_states:
                rows.append([pop.population, sel.generic_name, ms, sel.selection_basis, sel.combination_type])
    _write_sheet(ws, ["Population", "Comparator", "Member State", "Selection Basis", "Combination Type"], rows)

    # Sheet 5 - Agent 3 PICO Sets
    ws = wb.create_sheet("5. Agent3 PICO Sets")
    rows = []
    for pop in result.populations:
        for s in pop.assembly.pico_sets:
            rows.append([
                pop.population, s.pico_set_id, s.comparator, s.intervention, s.outcomes,
                s.member_states, len(s.member_states), s.tiers, s.rationale,
                s.combination_type, s.selection_basis, s.source_comparator_refs,
            ])
    _write_sheet(ws, ["Population", "PICO Set ID", "Comparator (C)", "Intervention (I)", "Outcomes (O)",
                       "Member States", "MS Count", "Tiers", "Rationale", "Combination Type",
                       "Selection Basis", "Source Comparator Refs"], rows)

    # Sheet 6 - Agent 4 Validation
    ws = wb.create_sheet("6. Agent4 Validation")
    rows = []
    for pop in result.populations:
        for r in pop.validation.results:
            rows.append([r.population, r.set_id, r.validation_result, r.failure_reason or ""])
    rows.append(["", "MEMBER STATE COUNT CHECK", f"total={result.populations[0].validation.member_state_count_check.total if result.populations else 27}",
                 str([p.validation.member_state_count_check.matches for p in result.populations])])
    _write_sheet(ws, ["Population", "Set ID", "Validation Result", "Failure Reason"], rows)

    # Sheet 7 - Evidence / Sources
    ws = wb.create_sheet("7. Evidence Sources")
    rows = []
    input_pop_by_name = {p.indication_disease: p for p in phase2_input.populations}
    for pop in result.populations:
        input_pop = input_pop_by_name.get(pop.population)
        comparator_by_name = {c.generic_name: c for c in (input_pop.comparators if input_pop else [])}
        for s in pop.assembly.pico_sets:
            for ref in (s.source_comparator_refs or [s.comparator]):
                comp = comparator_by_name.get(ref)
                if comp is None:
                    continue
                for req in comp.per_member_state:
                    rows.append([
                        pop.population, s.pico_set_id, comp.generic_name, req.member_state, req.tier,
                        req.source_reference, req.evidence_quote,
                    ])
    _write_sheet(ws, ["Population", "PICO Set ID", "Comparator", "Member State", "Tier",
                       "Source Reference", "Evidence Quote"], rows)

    # Sheet 8 - Not Identified Member States
    ws = wb.create_sheet("8. Not Identified MS")
    rows = []
    for pop in result.populations:
        for ms in pop.assembly.not_identified_states:
            rows.append([pop.population, ms])
    _write_sheet(ws, ["Population", "Not Identified Member State"], rows)

    return wb


def save_workbook(result: Phase2RunResult, phase2_input: Phase2Input, path: str) -> str:
    wb = build_workbook(result, phase2_input)
    wb.save(path)
    return path
