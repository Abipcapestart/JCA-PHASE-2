"""Word-document deliverable - the SME's actual named final output format
(architecture doc Section 7: "Word-document output (final deliverable, prose
spec only - no JSON given by SME): one column per PICO set, ordered
broadest-first; rows for Member States, Source Tier, Intervention (constant),
Outcomes (constant, efficacy/safety/QoL grouping preserved), Rationale (never
blank)."). This was never built in the first implementation pass (Excel/JSON
only) - this module closes that gap.

One table per population, columns = that population's PICO sets in the order
`SetAssemblyOutput.pico_sets` already carries (agent3_set_assembly.py sorts
broadest-first by Member State count - not re-sorted here).

`build_table_rows()` is the one place the row/column data is computed -
shared by the actual .docx builder below and by ui/app.py's on-screen
preview, so the on-screen preview and the downloaded Word document can never
drift apart into two different views of the same run.

Uses python-docx - the one dependency this module adds (see requirements.txt);
no other module in this package needs it.
"""

from __future__ import annotations

from typing import Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from schemas import OutcomeGroupLabel, Phase2Input, Phase2RunResult, PopulationInstance, PopulationResult

_GROUP_ORDER = [
    (OutcomeGroupLabel.EFFICACY, "Efficacy"),
    (OutcomeGroupLabel.SAFETY, "Safety"),
    (OutcomeGroupLabel.QUALITY_OF_LIFE_AND_SYMPTOMS, "Quality of Life & Symptoms"),
]


def _grouped_outcome_rows(population: Optional[PopulationInstance]) -> List[Tuple[str, str]]:
    """Outcomes are constant across every PICO set for a population (same
    list attached to each set by agent3_set_assembly.py) - grouped here by
    the SME's 3 named groups (efficacy/safety/QoL & symptoms), reading the
    group off Phase 2's *input* (`OutcomeInput.group`), since `Phase2PicoSet.
    outcomes` is just a flat name list with the grouping already collapsed
    out. An outcome not marked `listed` is included and annotated
    "(unlisted)" - per the SME's "listed or unlisted" input requirement,
    both belong in the deliverable, not just listed ones."""
    by_group: Dict[OutcomeGroupLabel, List[str]] = {g: [] for g, _ in _GROUP_ORDER}
    for o in (population.outcomes if population else []):
        text = o.concept if o.listed else f"{o.concept} (unlisted)"
        by_group.setdefault(o.group, []).append(text)
    rows = []
    for group, label in _GROUP_ORDER:
        items = by_group.get(group, [])
        rows.append((f"Outcomes (O) — {label}", "; ".join(items) if items else "—"))
    return rows


def build_table_rows(
    pop_result: PopulationResult, population: Optional[PopulationInstance],
) -> Tuple[List[str], List[Tuple[str, List[str]]]]:
    """Returns `(column_headers, rows)`: `column_headers` is one entry per
    PICO set (its comparator name, broadest-first); `rows` is a list of
    `(row_label, values)` where `values` has exactly `len(column_headers)`
    entries, one per PICO set, in the same order. `column_headers` is empty
    when the population produced no PICO sets - callers must check that
    before rendering a table."""
    pico_sets = pop_result.assembly.pico_sets
    column_headers = [s.comparator for s in pico_sets]
    if not pico_sets:
        return column_headers, []

    rows: List[Tuple[str, List[str]]] = [
        ("Member States", [", ".join(s.member_states) for s in pico_sets]),
        ("Source Tier", [", ".join(str(t) for t in sorted(set(s.tiers))) for s in pico_sets]),
        ("Intervention (I)", [s.intervention or "—" for s in pico_sets]),
    ]
    for label, text in _grouped_outcome_rows(population):
        rows.append((label, [text] * len(pico_sets)))
    # Never blank (SME requirement) - a failed generation call already
    # produces an explicit "[RATIONALE GENERATION FAILED...]" marker
    # (agent3_set_assembly.py), so this is never silently empty either way.
    rows.append(("Rationale", [s.rationale or "[MISSING - flagged for review]" for s in pico_sets]))
    return column_headers, rows


def _set_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text) if text else "")
    run.bold = bold
    run.font.size = Pt(10)


def _add_population_section(document: Document, pop_result: PopulationResult,
                             population: Optional[PopulationInstance]) -> None:
    document.add_heading(f"{pop_result.population} ({pop_result.population_label})", level=1)

    column_headers, rows = build_table_rows(pop_result, population)
    if not column_headers:
        document.add_paragraph("No PICO sets were produced for this population.")
        return

    n_cols = 1 + len(column_headers)
    table = document.add_table(rows=1 + len(rows), cols=n_cols)
    table.style = "Light Grid Accent 1"

    _set_cell(table.cell(0, 0), "PICO Set (Comparator, C)", bold=True)
    for i, name in enumerate(column_headers, start=1):
        _set_cell(table.cell(0, i), name, bold=True)

    for r, (label, values) in enumerate(rows, start=1):
        _set_cell(table.cell(r, 0), label, bold=True)
        for i, value in enumerate(values, start=1):
            _set_cell(table.cell(r, i), value)

    if pop_result.assembly.not_identified_states:
        document.add_paragraph()
        p = document.add_paragraph()
        p.add_run("Not-identified Member States (no comparator/standard of care found at scoping): ").bold = True
        p.add_run(", ".join(pop_result.assembly.not_identified_states))


def match_population_input(phase2_input: Phase2Input, pop_result: PopulationResult) -> Optional[PopulationInstance]:
    """Matches a `PopulationResult` (engineering-defined run output) back to
    the `PopulationInstance` (Phase 2's input) it was computed from, so
    input-only fields (outcome grouping, `listed`) can be attached to the
    output-only PICO set rows. Matches on (indication, population_label)
    first since two populations can share an indication (licensed vs ITT),
    falling back to indication alone if the label doesn't line up."""
    by_key = {(p.indication_disease, p.population_label): p for p in phase2_input.populations}
    by_indication = {p.indication_disease: p for p in phase2_input.populations}
    return by_key.get((pop_result.population, pop_result.population_label)) \
        or by_indication.get(pop_result.population)


def build_document(result: Phase2RunResult, phase2_input: Phase2Input) -> Document:
    document = Document()
    document.add_heading("JCA Phase 2 — PICOS Set Consolidation", level=0)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta_run = meta.add_run(
        f"Run ID: {result.run_id}   |   Model: {result.model or 'unknown'}   |   "
        f"Status: {result.overall_status}"
    )
    meta_run.italic = True
    meta_run.font.size = Pt(9)

    for pop_result in result.populations:
        population = match_population_input(phase2_input, pop_result)
        _add_population_section(document, pop_result, population)

    return document


def save_document(result: Phase2RunResult, phase2_input: Phase2Input, path: str) -> str:
    document = build_document(result, phase2_input)
    document.save(path)
    return path
