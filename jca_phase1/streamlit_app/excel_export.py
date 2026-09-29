"""
Excel export for a single full-workflow run record (see run_store.py for the
record shape). Two halves:

  1. The FINAL output's nested detail that's already in Phase1Output.to_dict()
     but wasn't being surfaced before: each comparator's/outcome's sources,
     evidence, and (for comparators) scope-adjudication reasoning -- exactly
     the "why is this included" + "sources" detail the production mock UI
     shows per item.
  2. The full per-stage debug trail (run["debug_capture"], populated by
     jca_phase1.orchestrator.run_phase1's debug_capture hook): every record
     extracted at A8, its grounding state at A9, its validation verdict at
     A10, and every candidate comparator's scope adjudication at A12 --
     including the ones that did NOT make the final comparator list. This is
     the "where did it get lost" trail an SME needs when a real comparator or
     outcome is missing from the final output.
"""

from __future__ import annotations

import io
import json
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DEBUG_HEADER_FILL = PatternFill("solid", fgColor="7C2D12")


def _write_table(ws, headers, rows, header_fill=HEADER_FILL):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = max(12, min(60, len(str(h)) + 4))


def _record_row(r: Dict[str, Any]) -> list:
    comp = r.get("comparator") or {}
    outc = r.get("outcome") or {}
    v = r.get("validation") or {}
    return [
        r.get("finding_id"), r.get("finding_type"), r.get("subject_drug"),
        r.get("member_state"), r.get("source_class"), r.get("tier"),
        r.get("source_url"), r.get("organization"),
        comp.get("as_stated") or outc.get("measure"),
        comp.get("role") if r.get("finding_type") == "comparator" else outc.get("result"),
        (r.get("evidence_quote") or "")[:300],
        r.get("grounded"), r.get("grounding_note"),
        v.get("verdict"), v.get("reason"), v.get("refetched"),
    ]


_RECORD_HEADERS = ["finding_id", "finding_type", "subject_drug", "member_state",
                   "source_class", "tier", "source_url", "organization",
                   "comparator_or_outcome", "role_or_result", "evidence_quote",
                   "grounded", "grounding_note", "validation_verdict",
                   "validation_reason", "refetched"]


def build_excel(run: Dict[str, Any]) -> bytes:
    wb = Workbook()
    out = run.get("output") or {}
    debug = run.get("debug_capture") or {}

    # ---- Run Summary ----------------------------------------------------
    ws = wb.active
    ws.title = "Run Summary"
    v = out.get("validation", {})
    _write_table(ws, ["field", "value"], [
        ["run_id", run.get("run_id")], ["status", run.get("status")],
        ["total_latency_s", run.get("total_latency_s")],
        ["total_cost_usd", run.get("total_cost_usd")],
        ["bedrock_cost_usd", run.get("bedrock_summary", {}).get("total_cost_usd")],
        ["tavily_cost_usd", run.get("tavily_summary", {}).get("total_cost_usd")],
        ["num_comparators", len(out.get("comparators", []))],
        ["records_extracted", v.get("records_extracted")],
        ["grounded", v.get("grounded")], ["grounding_failed", v.get("grounding_failed")],
        ["claim_validated", v.get("claim_validated")],
        ["claim_rejected", v.get("claim_rejected")],
        ["subject_drug_rejections", v.get("subject_drug_rejections")],
        ["role_rejections", v.get("role_rejections")],
        ["scope_in", v.get("scope_in")], ["scope_out", v.get("scope_out")],
        ["scope_uncertain", v.get("scope_uncertain")],
    ])

    ws = wb.create_sheet("Prompt Versions Used")
    _write_table(ws, ["prompt_id", "version"],
                [[k, v_] for k, v_ in run.get("prompt_versions_used", {}).items()])

    # ---- A1 input validation ----------------------------------------------
    ws = wb.create_sheet("A1 Input Validation")
    iv = out.get("input_validation", {}) or run.get("input_validation", {}) or {}
    rows = [["FIX", "; ".join(i.get("fields", [])), i.get("value_flagged", ""), i.get("explanation", "")]
           for i in iv.get("fix_items", [])]
    rows += [["CHECK", "; ".join(i.get("fields", [])), "", i.get("explanation", "")]
            for i in iv.get("check_items", [])]
    _write_table(ws, ["severity", "fields", "value_flagged", "explanation"], rows)

    # ---- A2 structured population / intervention ---------------------------
    ws = wb.create_sheet("A2 Population Fields")
    rows = []
    for p in out.get("populations", []):
        for fname, f in p.get("fields", {}).items():
            rows.append([p.get("population_id"), fname, f.get("value"), f.get("provenance"),
                        f.get("inference_basis")])
    _write_table(ws, ["population_id", "field", "value", "provenance", "inference_basis"], rows)

    ws = wb.create_sheet("A2 Intervention Fields")
    interv = out.get("intervention", {}) or {}
    rows = [[fname, f.get("value"), f.get("provenance"), f.get("inference_basis")]
           for fname, f in interv.get("fields", {}).items()]
    _write_table(ws, ["field", "value", "provenance", "inference_basis"], rows)

    # ---- A3 scope boundary --------------------------------------------------
    ws = wb.create_sheet("A3 Scope Boundary")
    rows = []
    for b in out.get("scope_boundaries", []):
        unbounded = set(b.get("unbounded_facets", []))
        for fname, f in b.get("facets", {}).items():
            is_bound = bool(f.get("value")) and f.get("provenance") in ("confirmed", "derived_from_label")
            rows.append([b.get("population_id"), fname, f.get("value"), f.get("provenance"),
                        f.get("must_match"), is_bound, fname in unbounded])
    _write_table(ws, ["population_id", "facet", "value", "provenance", "must_match",
                     "is_bound_filters_evidence", "explicitly_unbounded"], rows)

    # ---- A4 licensed indication ---------------------------------------------
    ws = wb.create_sheet("A4 Licensed Indication")
    lic = out.get("licensed_indication_record", {}) or {}
    _write_table(ws, ["field", "value"], [
        ["source", lic.get("source")], ["indication_text", lic.get("indication_text")],
        ["source_url", lic.get("source_url")], ["retrieved_at", lic.get("retrieved_at")],
        ["pivotal_trials", "; ".join(lic.get("pivotal_trials", []))],
        ["atc_code", lic.get("atc_code")], ["inn", lic.get("inn")], ["notes", lic.get("notes")],
    ])

    # ---- A5 therapeutic area(s) ---------------------------------------------
    ws = wb.create_sheet("A5 Therapeutic Areas")
    areas = out.get("therapeutic_areas", {}) or {}
    _write_table(ws, ["area", "rationale"],
                [[a, areas.get("rationales", {}).get(a, "")] for a in areas.get("areas", [])])
    ws.append([])
    ws.append(["method", areas.get("method")])
    ws.append(["needed_adjudication", areas.get("needed_adjudication")])

    # ---- Run manifest --------------------------------------------------------
    ws = wb.create_sheet("Run Manifest")
    m = out.get("run_manifest", {}) or {}
    workbook_info = m.get("source_workbook", {}) or {}
    _write_table(ws, ["field", "value"], [
        ["model", m.get("model")], ["code_version", m.get("code_version")],
        ["outcome_catalog_version", m.get("outcome_catalog_version")],
        ["therapeutic_areas_applied", "; ".join(m.get("therapeutic_areas_applied", []))],
        ["tier3_enabled", m.get("tier3_enabled")],
        ["source_workbook_path", workbook_info.get("path")],
        ["source_workbook_entries", workbook_info.get("entries")],
        ["source_workbook_unique_domains", workbook_info.get("unique_domains")],
        ["source_workbook_problems", json.dumps(workbook_info.get("problems", []))],
        ["leakage_guard", json.dumps(m.get("leakage_guard", {}))],
        ["providers", json.dumps(m.get("providers", {}))],
        ["started_at", m.get("started_at")], ["finished_at", m.get("finished_at")],
        ["duration_s", m.get("duration_s")],
        ["llm_usage", json.dumps(m.get("llm_usage", {}))],
        ["search_usage", json.dumps(m.get("search_usage", {}))],
    ])

    # ---- Validation: every excluded item, any stage --------------------------
    ws = wb.create_sheet("Validation Excluded Items")
    rows = [[e.get("value"), e.get("stage"), e.get("reason"), e.get("decisive_facet", ""),
            e.get("source_url", "")] for e in out.get("validation", {}).get("excluded", [])]
    _write_table(ws, ["value", "stage", "reason", "decisive_facet", "source_url"], rows)

    # ---- Final comparators + their full sources/evidence/adjudication ----
    ws = wb.create_sheet("Comparators")
    rows = []
    for c in out.get("comparators", []):
        rows.append([
            c.get("generic_name"), c.get("comparator_id"), "; ".join(c.get("brand_names", [])),
            c.get("inn"), c.get("atc_code"), c.get("class_or_mechanism"), c.get("class_source"),
            c.get("is_combination"), "; ".join(c.get("components", [])), c.get("role"),
            c.get("indication"), c.get("indication_scope_note"), c.get("line_of_therapy"),
            c.get("comparator_scenario"), c.get("retain_all_status"), c.get("recommendation_strength"),
            c.get("member_state_count"), "; ".join(c.get("member_states", [])),
            "; ".join(str(t) for t in c.get("tiers", [])), c.get("cross_tier_confirmed"),
            c.get("or_alternative"), c.get("general_evidence_flag"),
            c.get("rationale"), c.get("origin"), len(c.get("sources", [])), len(c.get("evidence", [])),
        ])
    _write_table(ws, ["generic_name", "comparator_id", "brand_names", "inn", "atc_code",
                     "class_or_mechanism", "class_source", "is_combination", "components", "role",
                     "indication", "indication_scope_note", "line_of_therapy", "comparator_scenario",
                     "retain_all_status", "recommendation_strength", "member_state_count",
                     "member_states", "tiers", "cross_tier_confirmed", "or_alternative",
                     "general_evidence_flag", "rationale", "origin",
                     "num_sources", "num_evidence"], rows)

    ws = wb.create_sheet("Comparator Sources")
    rows = []
    for c in out.get("comparators", []):
        for s in c.get("sources", []):
            rows.append([c.get("generic_name"), s.get("source_id"), s.get("tier"),
                        s.get("source_class"), s.get("organization"), s.get("display_name"),
                        s.get("url"), s.get("document_title"), s.get("document_date")])
    _write_table(ws, ["generic_name", "source_id", "tier", "source_class", "organization",
                     "display_name", "url", "document_title", "document_date"], rows)

    ws = wb.create_sheet("Comparator Evidence")
    rows = []
    for c in out.get("comparators", []):
        for e in c.get("evidence", []):
            rows.append([c.get("generic_name"), e.get("source_id"), (e.get("quote") or "")[:400],
                        e.get("locator"), e.get("member_state"), e.get("grounded"),
                        e.get("validation_verdict"), e.get("refetched")])
    _write_table(ws, ["generic_name", "source_id", "quote", "locator", "member_state",
                     "grounded", "validation_verdict", "refetched"], rows)

    ws = wb.create_sheet("Comparator Scope Adjudication")
    rows = []
    for c in out.get("comparators", []):
        adj = c.get("scope_adjudication") or {}
        rows.append([
            c.get("generic_name"), adj.get("verdict"), adj.get("decisive_facet"),
            adj.get("reason"), adj.get("adjudicator_version"),
            len(adj.get("evidence_for", [])), len(adj.get("evidence_against", [])),
        ])
    _write_table(ws, ["generic_name", "verdict", "decisive_facet", "reason",
                     "adjudicator_version", "num_evidence_for", "num_evidence_against"], rows)

    ws = wb.create_sheet("Per-State Verdicts")
    rows = []
    for c in out.get("comparators", []):
        for pv in c.get("per_member_state", []):
            rows.append([c.get("generic_name"), pv.get("member_state"), pv.get("verdict"),
                        pv.get("tier"), pv.get("source_url", ""),
                        (pv.get("evidence_quote") or "")[:300], pv.get("reason", "")])
    _write_table(ws, ["generic_name", "member_state", "verdict", "tier", "source_url",
                     "evidence_quote", "reason"], rows)

    # ---- Final outcomes + their full sources/evidence ---------------------
    outcomes = out.get("outcomes", {}) or {}
    ws = wb.create_sheet("Outcomes")
    rows = []
    for cat in ("clinical_effectiveness", "safety", "quality_of_life", "clinician_patient_reported"):
        for o in outcomes.get(cat, []):
            rows.append([cat, o.get("concept"), o.get("outcome_id"), o.get("category"),
                        o.get("catalog_id"), o.get("listed"), o.get("instrument"),
                        o.get("unit_of_measurement"), o.get("unit_disagreement_note"),
                        o.get("requirement_type"), o.get("coverage_status"),
                        "; ".join(str(t) for t in o.get("tiers", [])), o.get("cross_tier_confirmed"),
                        o.get("rationale"), o.get("origin"), "; ".join(o.get("aliases_merged", [])),
                        len(o.get("sources", [])), len(o.get("evidence", []))])
    _write_table(ws, ["bucket", "concept", "outcome_id", "category", "catalog_id", "listed",
                     "instrument", "unit_of_measurement", "unit_disagreement_note",
                     "requirement_type", "coverage_status", "tiers", "cross_tier_confirmed",
                     "rationale", "origin", "aliases_merged", "num_sources", "num_evidence"], rows)

    ws = wb.create_sheet("Outcome Sources")
    rows = []
    for cat in ("clinical_effectiveness", "safety", "quality_of_life", "clinician_patient_reported"):
        for o in outcomes.get(cat, []):
            for s in o.get("sources", []):
                rows.append([o.get("concept"), s.get("source_id"), s.get("tier"),
                            s.get("source_class"), s.get("organization"), s.get("display_name"),
                            s.get("url")])
    _write_table(ws, ["concept", "source_id", "tier", "source_class", "organization",
                     "display_name", "url"], rows)

    ws = wb.create_sheet("Outcome Evidence")
    rows = []
    for cat in ("clinical_effectiveness", "safety", "quality_of_life", "clinician_patient_reported"):
        for o in outcomes.get(cat, []):
            for e in o.get("evidence", []):
                rows.append([o.get("concept"), e.get("source_id"), (e.get("quote") or "")[:400],
                            e.get("member_state"), e.get("grounded"), e.get("validation_verdict")])
    _write_table(ws, ["concept", "source_id", "quote", "member_state", "grounded",
                     "validation_verdict"], rows)

    ws = wb.create_sheet("Catalog Coverage")
    _write_table(ws, ["catalog_id", "display_name", "category", "status", "matched_outcome_id"],
                [[cc.get("catalog_id"), cc.get("display_name"), cc.get("category"),
                  cc.get("status"), cc.get("matched_outcome_id", "")]
                 for cc in outcomes.get("catalog_coverage", [])])

    ws = wb.create_sheet("By Member State")
    _write_table(ws, ["member_state", "status", "comparators", "state_search_status"],
                [[e.get("member_state"), e.get("status"), "; ".join(e.get("comparators", [])),
                  e.get("state_search_status")] for e in out.get("by_member_state", [])])

    # ---- Cost / latency telemetry -----------------------------------------
    ws = wb.create_sheet("Bedrock Calls")
    _write_table(ws, ["stage", "prompt_id", "prompt_version", "input_tokens", "output_tokens",
                     "cost_usd", "latency_s", "attempts", "error"],
                [[c.get("stage"), c["prompt_id"], c["prompt_version"], c["input_tokens"],
                  c["output_tokens"], c["cost_usd"], c["latency_s"], c["attempts"], c["error"]]
                 for c in run.get("bedrock_calls", [])])

    ws = wb.create_sheet("Tavily Calls")
    _write_table(ws, ["op", "detail", "credits", "cost_usd", "latency_s", "attempts", "results", "error"],
                [[c["op"], c["detail"], c["credits"], c["cost_usd"], c["latency_s"],
                  c["attempts"], c["results"], c["error"]] for c in run.get("tavily_calls", [])])

    ws = wb.create_sheet("Stage Latency")
    _write_table(ws, ["stage", "latency_s"],
                [[stage, secs] for stage, secs in run.get("stage_latency_s", {}).items()])

    ws = wb.create_sheet("Errors and Retries")
    _write_table(ws, ["t_s", "level", "message"],
                [[e["t"], e["level"], e["message"]] for e in run.get("error_log", [])])

    # =========================================================================
    # DEBUG TRAIL -- every stage's raw intermediate state, including
    # candidates that did NOT survive to the final output. Only present if
    # this run was executed after the debug_capture hook was added; older
    # cached runs simply produce empty sheets here.
    # =========================================================================

    ws = wb.create_sheet("Debug - A6 Query Plan")
    plan = (debug.get("A6_query_plan") or {}).get("plan", [])
    _write_table(ws, ["query", "source_class", "member_state", "pass_type", "domains",
                     "max_urls", "language", "note"],
                [[p.get("query"), p.get("source_class"), p.get("member_state"),
                  p.get("pass_type"), "; ".join(p.get("domains", [])), p.get("max_urls"),
                  p.get("language"), p.get("note")] for p in plan], DEBUG_HEADER_FILL)

    ws = wb.create_sheet("Debug - A7 Attempts")
    attempts = (debug.get("A7_retrieval") or {}).get("attempts", [])
    _write_table(ws, ["member_state", "source_class", "attempted", "documents_retrieved",
                     "status", "detail"],
                [[a.get("member_state"), a.get("source_class"), a.get("attempted"),
                  a.get("documents_retrieved"), a.get("status"), a.get("detail")]
                 for a in attempts], DEBUG_HEADER_FILL)

    ws = wb.create_sheet("Debug - A7 Documents")
    docs = (debug.get("A7_retrieval") or {}).get("documents", [])
    _write_table(ws, ["url", "member_state", "source_class", "organization", "ok", "status",
                     "title", "text_preview"],
                [[d.get("url"), d.get("member_state"), d.get("source_class"),
                  d.get("organization"), d.get("ok"), d.get("status"), d.get("title"),
                  (d.get("text") or "")[:300]] for d in docs], DEBUG_HEADER_FILL)

    for key, title in [("A8_extraction", "Debug - A8 Extraction"),
                       ("A9_grounding", "Debug - A9 Grounding"),
                       ("A10_claim_validation", "Debug - A10 Validation"),
                       ("usable_after_A10", "Debug - Usable After A10")]:
        ws = wb.create_sheet(title)
        recs = debug.get(key) or []
        _write_table(ws, _RECORD_HEADERS, [_record_row(r) for r in recs], DEBUG_HEADER_FILL)

    ws = wb.create_sheet("Debug - A11 Identity")
    identities = debug.get("A11_identity") or {}
    _write_table(ws, ["as_stated_key", "inn", "atc_code", "class_mechanism", "class_source",
                     "is_combination", "components"],
                [[k, v_.get("inn"), v_.get("atc_code"), v_.get("class_mechanism"),
                  v_.get("class_source"), v_.get("is_combination"),
                  "; ".join(v_.get("components", []))] for k, v_ in identities.items()],
                DEBUG_HEADER_FILL)

    ws = wb.create_sheet("Debug - A12 All Candidates")
    all_adj = debug.get("A12_scope_adjudication") or {}
    rows = []
    for key, d in all_adj.items():
        comp = d.get("comparator", {})
        adj = d.get("adjudication", {})
        rows.append([
            key, comp.get("as_stated"), comp.get("role"), d.get("num_records"),
            adj.get("verdict"), adj.get("decisive_facet"), adj.get("reason"),
            len(adj.get("evidence_for", [])), len(adj.get("evidence_against", [])),
        ])
    _write_table(ws, ["candidate_key", "as_stated", "role", "num_records", "verdict",
                     "decisive_facet", "reason", "num_evidence_for", "num_evidence_against"],
                rows, DEBUG_HEADER_FILL)
    ws2 = wb.create_sheet("Debug - A12 Evid For-Against")
    rows2 = []
    for key, d in all_adj.items():
        comp = d.get("comparator", {})
        adj = d.get("adjudication", {})
        for side, items in (("for", adj.get("evidence_for", [])), ("against", adj.get("evidence_against", []))):
            for ec in items:
                rows2.append([key, comp.get("as_stated"), side, ec.get("source_id"),
                            ec.get("member_state"), ec.get("tier"), (ec.get("quote") or "")[:300],
                            "; ".join(ec.get("facets_matched", [])), ec.get("facet_conflict"),
                            ec.get("detail")])
    _write_table(ws2, ["candidate_key", "as_stated", "side", "source_id", "member_state", "tier",
                      "quote", "facets_matched", "facet_conflict", "detail"], rows2, DEBUG_HEADER_FILL)

    ws = wb.create_sheet("Debug - A16 Dropped")
    _write_table(ws, ["value", "reason"],
                [[d.get("value"), d.get("reason")] for d in (debug.get("A16_dropped_at_consolidation") or [])],
                DEBUG_HEADER_FILL)

    ws = wb.create_sheet("Debug - A15 Outcome Groups")
    rows = []
    for g in (debug.get("A15_outcome_groups") or []):
        for r in g.get("records", []):
            rows.append([g.get("concept"), g.get("catalog_id")] + _record_row(r))
    _write_table(ws, ["group_concept", "group_catalog_id"] + _RECORD_HEADERS, rows, DEBUG_HEADER_FILL)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
