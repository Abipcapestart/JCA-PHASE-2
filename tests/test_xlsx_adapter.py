"""Regression test: `comparator_scenario`/`retain_all_status` must be read
straight through from the "Comparators" sheet of a jca_phase1 xlsx export,
not hardcoded blank. A prior version of xlsx_adapter.py hardcoded
`comparator_scenario` to "" on the (stale) belief that the sheet had no such
column - it does (see jca_phase1/streamlit_app/excel_export.py's
`build_excel()`, which actually produces this file) - so every xlsx-sourced
comparator was silently losing all its per-Member-State rows regardless of
what scoping had actually recorded."""

import io

from openpyxl import Workbook

from xlsx_adapter import load_phase1_xlsx

_COMPARATOR_HEADERS = [
    "generic_name", "comparator_id", "brand_names", "inn", "atc_code",
    "class_or_mechanism", "class_source", "is_combination", "components", "role",
    "indication", "indication_scope_note", "line_of_therapy", "comparator_scenario",
    "retain_all_status", "recommendation_strength", "member_state_count",
    "member_states", "tiers", "cross_tier_confirmed", "or_alternative",
    "general_evidence_flag", "rationale", "origin", "num_sources", "num_evidence",
]

_VERDICT_HEADERS = ["generic_name", "member_state", "verdict", "tier", "source_url",
                    "evidence_quote", "reason"]


def _build_workbook(comparator_scenario: str, retain_all_status: str) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Run Summary"
    ws.append(["field", "value"])
    ws.append(["status", "complete"])

    ws = wb.create_sheet("Comparators")
    ws.append(_COMPARATOR_HEADERS)
    row = {h: "" for h in _COMPARATOR_HEADERS}
    row.update({
        "generic_name": "Gefitinib", "indication": "NSCLC",
        "comparator_scenario": comparator_scenario, "retain_all_status": retain_all_status,
        "member_states": "Denmark",
    })
    ws.append([row[h] for h in _COMPARATOR_HEADERS])

    ws = wb.create_sheet("Per-State Verdicts")
    ws.append(_VERDICT_HEADERS)
    ws.append(["Gefitinib", "Denmark", "standard_of_care", 1, "https://example.org", "quote", ""])

    wb.create_sheet("Outcomes").append(["bucket", "concept"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_comparator_scenario_read_through_from_sheet():
    buf = _build_workbook("at_least_one", "confirmed_droppable")
    output = load_phase1_xlsx(buf, product_name="TestDrug")
    comparator = output["comparators"][0]
    assert comparator["comparator_scenario"] == "at_least_one"
    assert comparator["retain_all_status"] == "confirmed_droppable"
    # the per-state row must survive into the adapted output, not be skipped
    assert comparator["per_member_state"][0]["member_state"] == "Denmark"


def test_comparator_scenario_blank_cell_stays_blank_not_guessed():
    buf = _build_workbook("", "")
    output = load_phase1_xlsx(buf, product_name="TestDrug")
    comparator = output["comparators"][0]
    assert comparator["comparator_scenario"] == ""
    assert comparator["retain_all_status"] == "unconfirmed"
